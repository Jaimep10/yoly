"""Professional Excel generator for Yoly reports."""
import sys
import os

# Add parent directory to path to import generador_excel
sys.path.insert(0, '/home/claude/yoly')

from generador_excel import crear_excel_profesional


def create_excel_profesional(wa_id, cuenta, gastos=None, cobro=None):
    """
    Crea un Excel profesional con 5 pestañas.

    Wrapper around generador_excel.crear_excel_profesional para mantener compatibilidad.

    Args:
        wa_id: Teléfono del usuario
        cuenta: Nombre de la cuenta
        gastos: Lista de gastos
        cobro: Datos de cobro/deuda (opcional)

    Returns:
        Workbook object (openpyxl)
    """
    if gastos is None:
        gastos = []
    if cobro is None:
        cobro = {}

    return crear_excel_profesional(wa_id, cuenta, gastos, cobro)


# ==================== EXCEL CON FÓRMULAS (reporte web) ====================

def _hoja_movimientos(wb, titulo, filas, color):
    """Hoja Fecha | Descripción | Monto | Categoría con el total como fórmula =SUM(...).
    Devuelve la celda del total (ej. 'C12') para que el RESUMEN la use."""
    from openpyxl.styles import Font, PatternFill, Alignment

    ws = wb.create_sheet(titulo)
    encabezado = ["Fecha", "Descripción", "Monto", "Categoría"]
    for col, texto in enumerate(encabezado, 1):
        celda = ws.cell(row=1, column=col, value=texto)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=color)
        celda.alignment = Alignment(horizontal="center")

    for fila, m in enumerate(filas, 2):
        ws.cell(row=fila, column=1, value=m.get('fecha') or '')
        ws.cell(row=fila, column=2, value=m.get('descripcion') or m.get('proveedor') or m.get('cliente') or '')
        ws.cell(row=fila, column=3, value=_numero(m.get('monto'))).number_format = '$#,##0.00'
        ws.cell(row=fila, column=4, value=m.get('categoria') or '')

    # Una fila vacía entre los datos y el total: así el rango nunca incluye el encabezado ni el total
    ultima = len(filas) + 2
    fila_total = ultima + 1
    ws.cell(row=fila_total, column=2, value="TOTAL").font = Font(bold=True)
    total = ws.cell(row=fila_total, column=3, value=f"=SUM(C2:C{ultima})")
    total.font = Font(bold=True)
    total.number_format = '$#,##0.00'

    for col, ancho in zip("ABCD", (12, 40, 14, 18)):
        ws.column_dimensions[col].width = ancho
    return f"C{fila_total}"


def _numero(valor):
    try:
        return round(float(str(valor).replace('$', '').replace(',', '').strip() or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def crear_excel_formulas(wa_id, cuenta, gastos=None, ingresos=None):
    """
    Excel con RESUMEN, GASTOS e INGRESOS donde Ingresos, Gastos y Saldo son fórmulas:
    si el cliente cambia un monto en Excel, los totales se recalculan solos.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    resumen = wb.active
    resumen.title = "RESUMEN"
    total_gastos = _hoja_movimientos(wb, "GASTOS", gastos or [], "D32F2F")
    total_ingresos = _hoja_movimientos(wb, "INGRESOS", ingresos or [], "2E7D32")

    resumen['A1'] = "YOLY ASISTENTE CONTABLE"
    resumen['A1'].font = Font(bold=True, size=14)
    resumen['A2'] = f"Teléfono: {wa_id}  ·  Cuenta: {cuenta}"
    resumen['A4'], resumen['B4'] = "Ingresos", f"=INGRESOS!{total_ingresos}"
    resumen['A5'], resumen['B5'] = "Gastos", f"=GASTOS!{total_gastos}"
    resumen['A6'], resumen['B6'] = "Saldo", "=B4-B5"
    for fila in (4, 5, 6):
        resumen[f'A{fila}'].font = Font(bold=True)
        resumen[f'B{fila}'].number_format = '$#,##0.00'
    resumen['B6'].font = Font(bold=True)
    resumen.column_dimensions['A'].width = 14
    resumen.column_dimensions['B'].width = 40
    return wb
