# Resumen por período ("resumen del 1 al 20 de junio") y transferencias que se guardan
# solas en su carpeta (sueldo, renta, deuda...).
import json
import os
from datetime import date
from unittest.mock import MagicMock

import pytest

import main
from agents import reporter as agent_reporter
import reportes

TEL = "whatsapp:+593991234567"
PHONE = main.normalizar_telefono(TEL)
HOY = date(2026, 10, 4)


@pytest.fixture
def entorno(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "DATA_DIR", str(tmp_path))
    guardar_real = main.guardar_memoria
    # Solo la memoria por cuenta (dentro de tmp_path); la global vive en /app/data
    monkeypatch.setattr(main, "guardar_memoria",
                        lambda m, phone_clean="", cuenta="principal":
                        guardar_real(m, phone_clean, cuenta) if phone_clean else None)
    monkeypatch.setattr(main, "memoria_usuarios", {})
    monkeypatch.setattr(main, "temp_gastos", {})
    monkeypatch.setattr(main, "temp_productos", {})
    monkeypatch.setattr(main, "descargar_media_twilio", lambda url: b"img")
    monkeypatch.setattr(main, "convertir_a_webp", lambda b: b"webp")
    monkeypatch.setattr(reportes, "hoy_fecha", lambda: HOY)
    claude = MagicMock()
    claude.messages.create.return_value.content = [MagicMock(text="¿Qué necesitas?")]
    monkeypatch.setattr(main, "client", claude)
    return tmp_path


def enviar(texto):
    salida = main.Salida()
    main.procesar_mensaje(texto, TEL, "https://yoly.test", salida)
    return "\n".join(salida.textos)


def foto(vision, texto=""):
    main.client.messages.create.return_value.content = [MagicMock(text=json.dumps(vision))]
    info = {}
    respuesta = main.procesar_foto_inteligente("https://media", TEL, texto, info)
    return respuesta, info


def transferencia(monto, concepto=None, direccion="enviada", fecha="2026-06-10", persona="Juan Pérez"):
    return {"tipo": "transferencia", "concepto": concepto, "direccion": direccion, "beneficiario": persona,
            "ordenante": persona, "banco": "Pichincha", "pagos": [{"fecha": fecha, "monto": monto}],
            "descripcion": "Transferencia"}


# ---------- Períodos ----------

@pytest.mark.parametrize("texto, inicio, fin", [
    ("resumen del 1 al 20 de junio", date(2026, 6, 1), date(2026, 6, 20)),
    ("resumen junio 1-20", date(2026, 6, 1), date(2026, 6, 20)),
    ("del 15 de mayo al 10 de junio", date(2026, 5, 15), date(2026, 6, 10)),
    ("1/6 al 20/6", date(2026, 6, 1), date(2026, 6, 20)),
    ("gastos de junio 2025", date(2025, 6, 1), date(2025, 6, 30)),
    ("resumen de noviembre", date(2025, 11, 1), date(2025, 11, 30)),  # noviembre aún no llega: el del año pasado
    ("semana pasada", date(2026, 9, 21), date(2026, 9, 27)),
    ("mes pasado", date(2026, 9, 1), date(2026, 9, 30)),
    ("primera quincena de junio", date(2026, 6, 1), date(2026, 6, 15)),
    ("15 de diciembre al 10 de enero", date(2025, 12, 15), date(2026, 1, 10)),
    ("1 al 31 de junio", date(2026, 6, 1), date(2026, 6, 30)),
])
def test_parsear_periodo(texto, inicio, fin):
    p = reportes.parsear_periodo(texto, HOY)
    assert (p["inicio"], p["fin"]) == (inicio, fin)


def test_sin_periodo():
    assert reportes.parsear_periodo("hola yoly", HOY) is None


# ---------- Resumen por WhatsApp ----------

def test_resumen_del_1_al_20_de_junio_solo_cuenta_esas_fechas(entorno):
    main.guardar_gastos(TEL, [
        {"fecha": "2026-06-02", "descripcion": "Supermaxi", "categoria": "comida", "monto": 50},
        {"fecha": "2026-06-20", "descripcion": "Arriendo", "categoria": "renta", "monto": 380},
        {"fecha": "2026-06-21", "descripcion": "Fuera del rango", "categoria": "comida", "monto": 999},
        {"fecha": "2026-05-31", "descripcion": "Fuera del rango", "categoria": "comida", "monto": 999},
        {"fecha": "2026-06-05", "descripcion": "Desglose de gastos", "categoria": "desglose", "monto": 100,
         "desglose_json": {"transporte": 30, "comida": 60}, "reserva": 10},
    ])
    with open(os.path.join(main.obtener_ruta_datos(TEL), "ingresos.json"), "w") as f:
        json.dump([{"fecha": "2026-06-15", "descripcion": "Venta", "monto": 1000}], f)

    texto = enviar("Resumen del 1 al 20 de junio")

    assert "Transacciones del 1 al 20 de junio" in texto
    assert "Supermaxi" in texto and "Arriendo" in texto and "Venta" in texto
    assert "Fuera del rango" not in texto
    assert "Ingresos: $1,000.00" in texto
    assert "Egresos: $530.00" in texto          # 50 + 380 + 30 + 60 + 10 de reserva
    assert "Balance: $470.00" in texto
    assert texto.endswith(agent_reporter.PREGUNTA_REPORTE)
    assert "/download/" not in texto           # nada de PDF/Excel hasta que conteste 1 o 2
    assert "/download/periodo/593991234567/pdf?desde=2026-06-01&hasta=2026-06-20" in enviar("1")
    enviar("Resumen del 1 al 20 de junio")
    assert "/download/periodo/593991234567/excel?desde=2026-06-01&hasta=2026-06-20" in enviar("2")
    main.client.messages.create.assert_not_called()


def test_resumen_sin_datos_lo_dice(entorno):
    assert enviar("resumen de junio") == agent_reporter.SIN_TRANSACCIONES


def test_balance_de_un_mes_no_se_queda_en_la_deuda(entorno):
    main.guardar_cobro(PHONE, {"tipo": "cobro_deuda", "cliente": "Alan", "deuda": 1000, "pagado": 300, "saldo": 700,
                               "pagos": [{"fecha": "2026-09-01", "monto": 300}]})
    main.guardar_gastos(TEL, [{"fecha": "2026-09-03", "descripcion": "Luz", "categoria": "servicios", "monto": 40}])
    assert "Transacciones del septiembre" in enviar("balance de septiembre")
    assert "Te falta: $700" in enviar("balance")


def test_registrar_gastos_con_montos_no_es_resumen(entorno):
    assert main.pedido_resumen_periodo("gastos de hoy: comida $20 y taxi $5") is None
    assert main.pedido_resumen_periodo("mis gastos de junio") is not None
    assert main.pedido_resumen_periodo("mándame el pdf del 1 al 20 de junio") is not None
    assert main.pedido_resumen_periodo("pdf") is None


# ---------- Transferencias a carpetas ----------

def carpeta(nombre):
    return reportes.cargar_carpetas(main.DATA_DIR, PHONE).get(nombre, [])


def test_transferencia_de_sueldo_va_sola_a_su_carpeta(entorno):
    respuesta, info = foto(transferencia(500, "Sueldo mensual"))

    assert info["tipo"] == "transferencia"
    assert "carpeta *Sueldos*" in respuesta and "$500.00" in respuesta and "Juan Pérez" in respuesta
    assert [t["monto"] for t in carpeta("sueldo")] == [500]
    assert os.path.exists(os.path.join(main.DATA_DIR, TEL, "sueldo", "transacciones.json"))
    gastos = main.cargar_gastos(TEL)
    assert gastos[-1]["categoria"] == "sueldo" and gastos[-1]["fecha"] == "2026-06-10"
    assert "sueldo: $500.00" in enviar("resumen de junio")


@pytest.mark.parametrize("concepto, esperada", [
    ("Pago arriendo octubre", "renta"), ("Abono préstamo", "deuda"), ("Pago luz", "servicios")])
def test_otros_conceptos(entorno, concepto, esperada):
    foto(transferencia(100, concepto))
    assert len(carpeta(esperada)) == 1


def test_lo_que_escribe_el_usuario_manda_sobre_el_concepto(entorno):
    respuesta, _ = foto(transferencia(380, "Transferencia"), texto="esto es la renta")
    assert "carpeta *Renta*" in respuesta


def test_sin_concepto_queda_por_revisar_y_se_puede_mover(entorno):
    respuesta, _ = foto(transferencia(250, None))
    assert "Por revisar" in respuesta and "es renta" in respuesta
    assert len(carpeta("por_revisar")) == 1

    texto = enviar("es deuda")

    assert "a *Deudas*" in texto
    assert carpeta("por_revisar") == [] and [t["monto"] for t in carpeta("deuda")] == [250]
    assert main.cargar_gastos(TEL)[-1]["categoria"] == "deuda"


def test_transferencia_recibida_es_ingreso(entorno):
    foto(transferencia(800, None, direccion="recibida", persona="Cliente ABC"))
    assert len(carpeta("ingreso")) == 1
    assert main.cargar_gastos(TEL) == []
    assert "Ingresos: $800.00" in enviar("resumen de junio")


def test_es_renta_sin_transferencias_no_hace_nada_raro(entorno):
    assert main.mover_transferencia(TEL, "renta", "https://yoly.test") is None


# ---------- Panel, Excel y PDF del período ----------

def test_panel_excel_y_pdf_del_periodo(entorno):
    foto(transferencia(500, "Sueldo"))
    main.guardar_gastos(TEL, main.cargar_gastos(TEL) + [
        {"fecha": "2026-06-12", "descripcion": "<b>Super</b>", "categoria": "comida", "monto": 45.5}])
    cliente = main.app.test_client()
    q = "desde=2026-06-01&hasta=2026-06-20"

    panel = cliente.get(f"/dashboard/{PHONE}/periodo?{q}")
    assert panel.status_code == 200
    html = panel.get_data(as_text=True)
    assert "$545.50" in html and "&lt;b&gt;Super&lt;/b&gt;" in html

    excel = cliente.get(f"/download/periodo/{PHONE}/excel?{q}")
    assert excel.status_code == 200 and excel.data[:2] == b"PK"
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(excel.data))
    assert wb.sheetnames == ["Resumen", "Movimientos", "Por semana"]
    assert wb["Resumen"]["B5"].value == 545.5

    pdf = cliente.get(f"/download/periodo/{PHONE}/pdf?{q}")
    assert pdf.status_code == 200 and pdf.data[:4] == b"%PDF"

    carpetas = cliente.get(f"/dashboard/{PHONE}/carpetas")
    assert carpetas.status_code == 200 and "Sueldos" in carpetas.get_data(as_text=True)
