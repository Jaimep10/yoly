# Yoly no inventa datos: solo guarda lo que leyó en la foto o lo que el usuario escribió,
# separa bien ingresos (dinero que entra) de gastos (dinero que sale) y, si no está claro, pregunta.
import io
import json
from datetime import date
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

import main
import reportes
from agents import vision

TEL = "whatsapp:+593991234567"
PHONE = main.normalizar_telefono(TEL)
HOY = date(2026, 10, 5)


@pytest.fixture
def entorno(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "DATA_DIR", str(tmp_path))
    monkeypatch.setattr(main, "guardar_memoria", lambda *a, **k: None)
    monkeypatch.setattr(main, "obtener_cuenta_activa", lambda *a, **k: "principal")
    monkeypatch.setattr(main, "memoria_usuarios", {})
    monkeypatch.setattr(main, "temp_gastos", {})
    monkeypatch.setattr(main, "temp_productos", {})
    monkeypatch.setattr(main, "descargar_media_twilio", lambda url: b"img")
    monkeypatch.setattr(main, "convertir_a_webp", lambda b: b"webp")
    monkeypatch.setattr(main, "fecha_exif", lambda b: None)
    monkeypatch.setattr(reportes, "hoy_fecha", lambda: HOY)
    claude = MagicMock()
    claude.messages.create.return_value.content = [MagicMock(text="Hola")]
    monkeypatch.setattr(main, "client", claude)
    return tmp_path


def enviar(texto, tel=TEL):
    salida = main.Salida()
    main.procesar_mensaje(texto, tel, "https://yoly.test", salida)
    return "\n".join(salida.textos)


def foto(vision_json, texto=""):
    main.client.messages.create.return_value.content = [MagicMock(text=json.dumps(vision_json))]
    return main.procesar_fotos_whatsapp(["https://media"], TEL, texto, {}, cuenta="principal")


def guardados():
    gastos, ingresos = reportes.cargar_movimientos(main.DATA_DIR, TEL)
    return [g["monto"] for g in gastos], [i["monto"] for i in ingresos]


def resumen_mes():
    periodo = reportes.parsear_periodo("este mes", HOY)
    return main.resumen_periodo_usuario(TEL, periodo["inicio"], periodo["fin"], "principal")[0]


def transferencia(monto, direccion, concepto=None):
    return {"tipo": "transferencia", "concepto": concepto, "direccion": direccion, "beneficiario": "Ana",
            "ordenante": "Luis", "pagos": [{"fecha": "2026-10-02", "monto": monto}]}


# ---------- El Ojo ----------

def test_prompt_del_ojo_no_tiene_valores_de_ejemplo_para_copiar():
    # Con un ejemplo con números reales, Vision a veces devolvía esos números cuando no leía la foto
    for ejemplo in ("Maria Cristina", "3000", "Supermaxi", "Quito", "Arroz blanco", "Balu"):
        assert ejemplo not in vision.PROMPT


# ---------- Fotos: ingreso vs gasto ----------

def test_foto_de_ingreso_queda_solo_como_ingreso(entorno):
    respuesta = foto(transferencia(500, "recibida"))
    assert "recibiste" in respuesta
    assert guardados() == ([], [500])
    r = resumen_mes()
    assert (r["total_ingresos"], r["total_gastos"]) == (500, 0)


def test_factura_de_500_es_un_gasto_de_exactamente_500(entorno):
    respuesta = foto({"tipo": "factura_compra", "tienda": "Ferretería", "descripcion": "Materiales",
                      "pagos": [{"fecha": "2026-10-01", "monto": "500.00"}], "articulos": []})
    assert "$500" in respuesta
    assert guardados() == ([500.0], [])
    r = resumen_mes()
    assert (r["total_ingresos"], r["total_gastos"], r["balance"]) == (0, 500, -500)


def test_transferencia_sin_saber_si_entro_o_salio_no_se_suma_y_pregunta(entorno):
    respuesta = foto(transferencia(80, None, "Pago"))
    assert "no se ve si este dinero ENTRÓ o SALIÓ" in respuesta
    assert "enviaste" not in respuesta
    r = resumen_mes()
    assert (r["total_ingresos"], r["total_gastos"], r["sin_direccion"]) == (0, 0, 1)
    assert "no sé si fueron ingreso o gasto" in reportes.texto_resumen(r, reportes.parsear_periodo("este mes", HOY))

    enviar("es ingreso")
    r = resumen_mes()
    assert (r["total_ingresos"], r["total_gastos"], r["sin_direccion"]) == (80, 0, 0)


def test_lo_que_escribe_el_usuario_decide_si_entro_o_salio(entorno):
    foto(transferencia(200, None), texto="me pagaron por el trabajo")
    assert guardados() == ([], [200])


def test_factura_que_el_usuario_dice_que_es_venta_es_ingreso(entorno):
    respuesta = foto({"tipo": "otro", "cliente": "Cliente Pérez", "pagos": [{"fecha": "2026-10-03", "monto": 120}]},
                     texto="esta es una venta que me pagaron")
    assert "INGRESO" in respuesta
    assert guardados() == ([], [120])


# ---------- Mensajes escritos ----------

@pytest.mark.parametrize("texto, tipo, monto, categoria, fecha", [
    ("me pagaron 300 por una venta", "ingreso", 300, "ventas", "2026-10-05"),
    ("recibí $1.200 el 3 de octubre", "ingreso", 1200, "ingreso", "2026-10-03"),
    ("Gasté 45,50 ayer", "gasto", 45.5, "otro", "2026-10-04"),
    ("pagué 380 de renta", "gasto", 380, "renta", "2026-10-05"),
    ("compré 20 de gasolina 02/10", "gasto", 20, "otro", "2026-10-02"),
])
def test_movimiento_escrito(texto, tipo, monto, categoria, fecha):
    m = reportes.movimiento_de_texto(texto, HOY)
    assert (m["tipo"], m["monto"], m["categoria"], m["fecha"]) == (tipo, monto, categoria, fecha)


@pytest.mark.parametrize("texto", [
    "pagué 500 de renta y 200 de luz",   # dos montos
    "me pagaron 300 y pagué 100",          # entra y sale a la vez
    "vendí 3 tortas a 15 cada una",        # no se sabe el total
    "gano 1200 al mes",                    # presupuesto, no un movimiento
    "pagué la luz",                        # sin monto
    "resumen de gastos",
])
def test_si_no_esta_claro_no_se_adivina(texto):
    assert reportes.movimiento_de_texto(texto, HOY) is None


def test_mensaje_de_ingreso_y_de_gasto_se_guardan_donde_van(entorno):
    assert "INGRESO de $300.00" in enviar("me pagaron 300 por una venta")
    assert "GASTO de $50.00" in enviar("pagué 50 de luz")
    assert guardados() == ([50], [300])
    main.client.messages.create.assert_not_called()


def test_mensaje_confuso_pregunta_y_no_guarda_ni_llama_a_claude(entorno):
    respuesta = enviar("pagué 500 de renta y 200 de luz")
    assert "un movimiento por mensaje" in respuesta
    assert guardados() == ([], [])
    main.client.messages.create.assert_not_called()


# ---------- Reportes y Excel ----------

def test_informe_de_gastos_usa_los_datos_reales(entorno):
    enviar("gasté 45.50 en comida")
    respuesta = enviar("informe de gastos")
    assert "$45.50" in respuesta
    assert "150" not in respuesta and "Alimentos" not in respuesta


def test_excel_sin_datos_no_inventa_ni_cae_a_claude(entorno):
    respuesta = enviar("excel")
    assert "no armé el Excel" in respuesta
    main.client.messages.create.assert_not_called()


def test_excel_tiene_ingresos_gastos_y_balance(entorno):
    enviar("me pagaron 300 por una venta")
    enviar("pagué 50 de luz")
    respuesta = enviar("mándame el excel")
    assert "1 ingresos y 1 gastos" in respuesta
    url = respuesta.split("https://yoly.test", 1)[1].split()[0]
    excel = main.app.test_client().get(url)
    assert excel.status_code == 200
    hoja = load_workbook(io.BytesIO(excel.data))["Resumen"]
    assert [hoja["B4"].value, hoja["B5"].value, hoja["B6"].value] == [300, 50, 250]

    # La descarga del panel también trae ingresos y gastos (antes solo gastos)
    panel = main.app.test_client().get(f"/download/excel/{PHONE}")
    assert load_workbook(io.BytesIO(panel.data))["Resumen"]["B4"].value == 300


def test_respuesta_libre_tiene_prohibido_dar_numeros(entorno):
    enviar("hola yoly")
    sistema = main.client.messages.create.call_args.kwargs["system"]
    assert sistema == main.SISTEMA_CHARLA
    assert "NUNCA des montos" in sistema and "NUNCA digas que guardaste" in sistema


def test_analisis_financiero_no_mezcla_usuarios(entorno):
    main.client.messages.create.return_value.content = [
        MagicMock(text='{"ingresos": 1000, "gastos": {"renta": 300}}')]
    main.procesar_analisis_financiero("gano 1000 y pago 300 de renta", TEL)
    main.client.messages.create.return_value.content = [MagicMock(text='{"ingresos": 0, "gastos": {}}')]
    _, ingresos, gastos, _ = main.procesar_analisis_financiero("hola", "whatsapp:+593990000000")
    assert (ingresos, gastos) == (0, 0)


def test_metas_son_de_cada_usuario_y_no_hay_metas_de_prueba(entorno):
    # Antes un goals.json del repo mostraba a todos "Pagar todas las deudas $200/$1300", "Fondo de emergencia"...
    respuesta = enviar("ver metas")
    assert "No tienes metas" in respuesta and "1300" not in respuesta and "Fondo de emergencia" not in respuesta
    main.registrar_meta("Moto", 2000, "2027-06-01", "savings", 100, telefono=TEL)
    assert "Moto" in enviar("ver metas")
    assert "No tienes metas" in enviar("ver metas", tel="whatsapp:+593990000000")
