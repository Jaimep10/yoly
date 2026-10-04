# Los 4 agentes por separado (Portero, Ojo, Calculadora, Contadora) y el flujo completo de app.py.
import io
import json
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from agents import calculator as agent_calculator
from agents import classifier as agent_classifier
from agents import reporter as agent_reporter
from agents import vision as agent_vision
import main

TEL = "whatsapp:+593991234567"
PHONE = main.normalizar_telefono(TEL)

PAGOS = [
    {"fecha": "2026-01-03", "monto": 200, "metodo": "efectivo"},
    {"fecha": "2026-01-15", "monto": 120, "metodo": "T"},
    {"fecha": "2026-02-01", "monto": 300, "metodo": "transf"},
    {"fecha": "2026-02-15", "monto": 160, "metodo": ""},
    {"fecha": "2026-03-01", "monto": 500, "metodo": "cheque"},
    {"fecha": "2026-03-15", "monto": 250, "metodo": "zelle"},
    {"fecha": "2026-04-01", "monto": 100, "metodo": "efec"},
    {"fecha": "2026-04-15", "monto": 150, "metodo": "dep"},
]
LIBRETA = {"cliente": "Maria Cristina", "deuda": 3000, "pagos": PAGOS, "tipo": "libreta_cobros"}


def respuesta_claude(datos):
    r = MagicMock()
    r.content = [MagicMock(text=datos if isinstance(datos, str) else json.dumps(datos))]
    return r


# ---------- Calculadora ----------

def test_calculadora_suma_en_python():
    cobro = agent_calculator.calcular_totales(LIBRETA)
    assert cobro["deuda"] == 3000 and cobro["pagado"] == 1780 and cobro["saldo"] == 1220
    assert cobro["valido"] and not cobro["duplicado"]
    assert [f["saldo"] for f in cobro["filas"]][:2] == [2800, 2680]
    assert [p["metodo"] for p in cobro["pagos"]][:4] == ["efectivo", "transferencia", "transferencia", "no especificado"]


def test_calculadora_detecta_duplicado_por_hash():
    guardado = agent_calculator.calcular_totales(LIBRETA)
    otra_vez = agent_calculator.calcular_totales(dict(LIBRETA, cliente="maria cristina "), guardado)
    assert otra_vez["hash"] == guardado["hash"] and otra_vez["duplicado"]
    con_pago_nuevo = dict(LIBRETA, pagos=PAGOS + [{"fecha": "2026-05-01", "monto": 100}])
    assert not agent_calculator.calcular_totales(con_pago_nuevo, guardado)["duplicado"]


def test_calculadora_combina_paginas_sin_contar_dos_veces():
    pagina1 = {"cliente": "Maria Cristina", "deuda": 3000, "pagos": PAGOS[:5]}
    pagina2 = {"cliente": None, "pagos": PAGOS[4:]}  # el pago de 500 sale en las dos fotos
    cobro = agent_calculator.calcular_totales([pagina1, pagina2])
    assert cobro["cliente"] == "Maria Cristina" and len(cobro["pagos"]) == 8
    assert cobro["pagado"] == 1780 and cobro["saldo"] == 1220


# ---------- Portero ----------

def test_portero_clasifica_y_cuenta():
    claude = MagicMock()
    claude.messages.create.return_value = respuesta_claude({"tipos": ["libreta_cobros", "factura"], "confianza": 0.9})
    r = agent_classifier.clasificar_documentos([b"a", b"b"], claude, "m")
    assert r["cantidad"] == 2 and r["confianza"] == 0.9
    assert [i["tipo"] for i in r["imagenes"]] == ["libretita_deuda", "facturas"]


def test_portero_para_fotos_repetidas_sin_llamar_a_claude():
    claude = MagicMock()
    vista = agent_classifier.huella_imagen(b"a")
    r = agent_classifier.clasificar_documentos([b"a", b"a"], claude, "m", vistas=[vista])
    assert r["tipo"] == "duplicado" and r["cantidad"] == 2
    claude.messages.create.assert_not_called()


def test_portero_salta_la_repetida_dentro_del_mismo_envio():
    claude = MagicMock()
    claude.messages.create.return_value = respuesta_claude({"tipos": ["libretita_deuda"]})
    r = agent_classifier.clasificar_documentos([b"a", b"a"], claude, "m")
    assert r["tipo"] == "libretita_deuda"
    assert [i["duplicada"] for i in r["imagenes"]] == [False, True]


def test_portero_si_falla_deja_decidir_al_ojo():
    claude = MagicMock()
    claude.messages.create.side_effect = RuntimeError("sin red")
    assert agent_classifier.clasificar_documentos([b"a"], claude, "m")["tipo"] == "otro"


# ---------- Ojo ----------

def test_ojo_devuelve_json_con_deuda():
    claude = MagicMock()
    claude.messages.create.return_value = respuesta_claude(
        "```json\n" + json.dumps({"cliente": "Ana", "deuda_total": 500, "pagos": [{"monto": 100}]}) + "\n```")
    datos = agent_vision.extraer_json(b"img", "libretita_deuda", claude, "m")
    assert datos["deuda"] == 500 and "deuda_total" not in datos
    prompt = claude.messages.create.call_args.kwargs["messages"][0]["content"][1]["text"]
    assert "libretita" in prompt and "No sumes" in prompt


def test_ojo_sin_json_devuelve_none():
    claude = MagicMock()
    claude.messages.create.return_value = respuesta_claude("no veo nada")
    assert agent_vision.extraer_json(b"img", None, claude, "m") is None


# ---------- Contadora ----------

def test_contadora_arma_mensaje_tabla_pdf_y_excel():
    cobro = agent_calculator.calcular_totales(LIBRETA)
    guardados = []
    reporte = agent_reporter.generar_reporte(cobro, PHONE, guardar=lambda p, d: guardados.append((p, d)))
    msg = reporte["mensaje_wa"]
    assert "Leí 8 pagos de Maria Cristina" in msg
    assert "Fecha     Monto    Método        Saldo" in msg and "03-01-26" in msg
    assert "Pagado: $1,780" in msg and "Te falta: $1,220" in msg and msg.endswith(agent_reporter.PREGUNTA_TABLA)
    assert reporte["dashboard_cards"]["saldo"] == 1220 and reporte["dashboard_cards"]["numero_pagos"] == 8
    assert reporte["pdf_bytes"].startswith(b"%PDF")
    wb = load_workbook(io.BytesIO(reporte["excel_bytes"]))
    assert wb.sheetnames == ["Pagos", "Resumen"]
    (phone, registro), = guardados
    assert phone == PHONE and registro["deuda"] == 3000 and registro["pagado"] == 1780 and registro["hash"]


def test_contadora_no_guarda_duplicados():
    cobro = agent_calculator.calcular_totales(LIBRETA)
    cobro["duplicado"] = True
    guardados = []
    reporte = agent_reporter.generar_reporte(cobro, PHONE, guardar=lambda p, d: guardados.append(d),
                                             incluir_archivos=False)
    assert not guardados and "ya la tenía guardada" in reporte["mensaje_wa"]


# ---------- Flujo completo (orchestrator.py) ----------

@pytest.fixture
def entorno(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "DATA_DIR", str(tmp_path))
    monkeypatch.setattr(main, "guardar_memoria", lambda m: None)
    monkeypatch.setattr(main, "memoria_usuarios", {})
    fotos = {"https://m/1": b"pagina1", "https://m/2": b"pagina2"}
    monkeypatch.setattr(main, "descargar_media_twilio", lambda url: fotos[url])
    monkeypatch.setattr(main, "convertir_a_webp", lambda b: b)
    claude = MagicMock()
    monkeypatch.setattr(main, "client", claude)
    return claude


def test_dos_fotos_de_libretita_hacen_un_solo_registro(entorno):
    entorno.messages.create.side_effect = [
        respuesta_claude({"tipos": ["libretita_deuda", "libretita_deuda"], "confianza": 0.95}),
        respuesta_claude({"cliente": "Maria Cristina", "deuda": 3000, "pagos": PAGOS[:4], "tipo": "libreta_cobros"}),
        respuesta_claude({"cliente": "Maria Cristina", "pagos": PAGOS[4:], "tipo": "libreta_cobros"}),
    ]
    info = {}
    texto = main.procesar_fotos_whatsapp(["https://m/1", "https://m/2"], TEL, "", info)
    assert info["clasificacion"]["cantidad"] == 2
    assert "Leí 8 pagos de Maria Cristina" in texto and "Te falta: $1,220" in texto
    assert texto.endswith(main.PREGUNTA_TABLA)
    cobro = main.obtener_cobro(PHONE)
    assert cobro["deuda"] == 3000 and cobro["pagado"] == 1780 and cobro["saldo"] == 1220
    assert main.obtener_ultima_pregunta(PHONE) == "dashboard"

    # Las mismas fotos otra vez: el Portero las para y no se llama a Vision
    entorno.messages.create.reset_mock()
    assert "ya la había procesado" in main.procesar_fotos_whatsapp(["https://m/1", "https://m/2"], TEL)
    entorno.messages.create.assert_not_called()


def test_webhook_manda_todas_las_fotos(entorno, monkeypatch):
    recibidas, enviados = [], []
    monkeypatch.setattr(main, "procesar_fotos_whatsapp",
                        lambda urls, tel, texto="", info=None: recibidas.append(urls) or "ok")
    monkeypatch.setattr(main, "enviar_por_twilio", lambda to, textos: enviados.append(textos))
    monkeypatch.setattr(main.lotes_fotos, "espera", 0.3)
    data = {"From": TEL, "Body": "", "NumMedia": "2", "MediaUrl0": "https://m/1", "MediaContentType0": "image/jpeg",
            "MediaUrl1": "https://m/2", "MediaContentType1": "image/jpeg"}
    with main.app.test_client() as c:
        r = c.post("/whatsapp", data=data)
    assert r.status_code == 200 and main.MENSAJE_FOTOS in r.get_data(as_text=True)
    assert esperar(lambda: enviados)
    assert recibidas == [["https://m/1", "https://m/2"]] and enviados == [["ok"]]


def esperar(condicion, segundos=5):
    import time
    fin = time.monotonic() + segundos
    while time.monotonic() < fin:
        if condicion():
            return True
        time.sleep(0.05)
    return False


def test_grupo_de_fotos_en_webhooks_separados_es_un_solo_lote(entorno, monkeypatch):
    """Twilio manda cada foto de un grupo en su propio webhook: se juntan y se procesan una sola vez."""
    recibidas, enviados = [], []
    monkeypatch.setattr(main, "procesar_fotos_whatsapp",
                        lambda urls, tel, texto="", info=None: recibidas.append(list(urls)) or "listo")
    monkeypatch.setattr(main, "enviar_por_twilio", lambda to, textos: enviados.append(textos))
    monkeypatch.setattr(main.lotes_fotos, "espera", 0.4)
    respuestas = []
    with main.app.test_client() as c:
        for i in range(7):
            data = {"From": TEL, "Body": "", "NumMedia": "1", "MediaUrl0": f"https://m/{i}",
                    "MediaContentType0": "image/jpeg"}
            respuestas.append(c.post("/whatsapp", data=data).get_data(as_text=True))
    # Solo el primer webhook dice "recibido"; los demás contestan vacío al instante
    assert main.MENSAJE_FOTOS in respuestas[0] and all(main.MENSAJE_FOTOS not in r for r in respuestas[1:])
    assert esperar(lambda: enviados)
    assert recibidas == [[f"https://m/{i}" for i in range(5)]]
    assert "leo máximo 5" in enviados[0][0] and "otras 2" in enviados[0][0] and enviados[0][1] == "listo"


def test_una_foto_que_no_contesta_no_traba_las_demas(monkeypatch):
    import time
    import orchestrator
    monkeypatch.setattr(orchestrator, "TIMEOUT_FOTO", 0.3)
    claude = MagicMock()
    gastos = []

    def crear(**kwargs):
        texto = kwargs["messages"][0]["content"][-1]["text"]
        if "portero" in texto:
            return respuesta_claude({"tipos": ["facturas"] * 3})
        datos = kwargs["messages"][0]["content"][0]["source"]["data"]
        if datos == main.base64.standard_b64encode(b"lenta").decode():
            time.sleep(2)   # Vision colgado con esta foto
        return respuesta_claude({"cliente": "Tienda", "pagos": [{"fecha": "2026-01-03", "monto": 10}]})

    claude.messages.create.side_effect = crear
    ctx = main.Contexto(
        cliente=claude, modelo="m", phone_clean=PHONE, guardar_imagen=lambda img: "/tmp/foto.webp",
        guardar_cobro=None, cobro_actual=lambda: None,
        guardar_gasto=lambda v, p, t, r: gastos.append(t) or f"Leí ${t}", guardar_transferencia=None,
        huellas_vistas=lambda: [], marcar_vistas=lambda h: None, marcar_pregunta_tabla=lambda: None)
    inicio = time.monotonic()
    texto = orchestrator.OrquestadorYoly().handle_whatsapp(PHONE, [b"a", b"lenta", b"c"], "", ctx=ctx)
    assert time.monotonic() - inicio < 1.5
    assert gastos == [10, 10] and "No pude leer la foto 2" in texto


def test_falla_una_descarga_sigue_con_las_otras(entorno, monkeypatch):
    monkeypatch.setattr(main, "descargar_media_twilio", lambda url: None if url.endswith("/2") else b"x" + url.encode())
    entorno.messages.create.side_effect = [
        respuesta_claude({"tipos": ["facturas"]}),
        respuesta_claude({"cliente": "Tienda", "pagos": [{"fecha": "2026-01-03", "monto": 25}]}),
    ]
    texto = main.procesar_fotos_whatsapp(["https://m/1", "https://m/2"], TEL)
    assert "No pude descargar la foto 2" in texto and "Leí Tienda: $25" in texto


def test_numero_twilio_con_prefijo_whatsapp(monkeypatch):
    monkeypatch.setenv("TWILIO_WHATSAPP_NUMBER", "+15163869020")
    assert main.numero_whatsapp_twilio() == "whatsapp:+15163869020"
    monkeypatch.setenv("TWILIO_WHATSAPP_NUMBER", "whatsapp:+15163869020")
    assert main.numero_whatsapp_twilio() == "whatsapp:+15163869020"


def test_descargas_pdf_y_excel_del_cobro(entorno):
    main.guardar_cobro(PHONE, dict(LIBRETA, tipo="cobro_deuda"))
    with main.app.test_client() as c:
        pdf = c.get(f"/download/pdf/{PHONE}")
        excel = c.get(f"/download/excel/{PHONE}")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")
    assert excel.status_code == 200
    assert load_workbook(io.BytesIO(excel.data))["Resumen"]["B5"].value == 1220


def test_orquestador_jefe_usa_los_4_agentes():
    from orchestrator import OrquestadorYoly
    claude = MagicMock()
    claude.messages.create.side_effect = [respuesta_claude({"tipos": ["libretita_deuda"]}), respuesta_claude(LIBRETA)]
    guardados, vistas = [], []
    ctx = main.Contexto(
        cliente=claude, modelo="m", phone_clean=PHONE, guardar_imagen=lambda img: "/tmp/foto.webp",
        guardar_cobro=lambda p, d: guardados.append(d), cobro_actual=lambda: None,
        guardar_gasto=None, guardar_transferencia=None, huellas_vistas=lambda: vistas,
        marcar_vistas=vistas.extend, marcar_pregunta_tabla=lambda: None)
    jefe = OrquestadorYoly()
    assert {type(a).__name__ for a in (jefe.portero, jefe.ojo, jefe.calculadora, jefe.contadora)} == \
        {"Portero", "Ojo", "Calculadora", "Contadora"}
    texto = jefe.handle_whatsapp(PHONE, [b"foto"], "", ctx=ctx)
    assert "Te falta: $1,220" in texto
    assert guardados[0]["deuda"] == 3000 and guardados[0]["pagado"] == 1780 and len(vistas) == 1
    assert jefe.handle_whatsapp(PHONE, [b"foto"], "", ctx=ctx) == "👍 Esa foto ya la había procesado, no la guardé otra vez."
