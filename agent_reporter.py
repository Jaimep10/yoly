# Agente 4 - "La Contadora": con las cuentas ya hechas por la Calculadora arma todo lo que ve el usuario:
# mensaje de WhatsApp (con tabla Fecha | Monto | Método | Saldo), tarjetas del dashboard,
# PDF con la tabla, Excel con hojas Pagos y Resumen, y guarda el registro en memoria_global.json.
import io
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

PREGUNTA_TABLA = "¿Te mando tabla al dashboard?\nResponde con el número:\n1️⃣ Sí\n2️⃣ No"
MAX_FILAS_WHATSAPP = 15


def titulo_cobro(cobro):
    return f"Estado de Cuenta {cobro['cliente']} - Deuda ${cobro['deuda']:,.0f} - Saldo ${cobro['saldo']:,.0f}"


def metodo_corto(metodo):
    return "—" if metodo == 'no especificado' else metodo.capitalize()


def tabla_whatsapp(cobro, max_filas=MAX_FILAS_WHATSAPP):
    """Tabla en texto monoespaciado para WhatsApp: Fecha | Monto | Método | Saldo."""
    con_saldo = cobro['deuda'] > 0
    filas = cobro['filas']
    lineas = ["Fecha     Monto    Método" + ("        Saldo" if con_saldo else "")]
    for f in filas[:max_filas]:
        linea = f"{f['fecha_corta']:<9} ${f['monto']:<7,.0f} {metodo_corto(f['metodo']):<13}"
        if con_saldo:
            linea += f" ${f['saldo']:,.0f}"
        lineas.append(linea.rstrip())
    texto = "```\n" + "\n".join(lineas) + "\n```"
    if len(filas) > max_filas:
        texto += f"\n…y {len(filas) - max_filas} pagos más (están todos en el dashboard)."
    return texto


def mensaje_whatsapp(cobro, duplicado=False):
    if duplicado:
        cabecera = f"👍 Esta libretita de {cobro['cliente']} ya la tenía guardada, no la sumé otra vez."
    else:
        cabecera = f"Leí {len(cobro['pagos'])} pagos de {cobro['cliente']}:"
    partes = [cabecera, tabla_whatsapp(cobro), f"Pagado: ${cobro['pagado']:,.0f}"]
    if cobro['deuda']:
        partes[-1] = f"Deuda original: ${cobro['deuda']:,.0f}\nPagado: ${cobro['pagado']:,.0f}\nTe falta: ${cobro['saldo']:,.0f}"
    return "\n".join(partes) + "\n\n" + PREGUNTA_TABLA


def dashboard_cards(cobro):
    """Lo que muestran las tarjetas del dashboard."""
    return {
        "cliente": cobro['cliente'],
        "deuda": cobro['deuda'],
        "pagado": cobro['pagado'],
        "saldo": cobro['saldo'],
        "numero_pagos": len(cobro['pagos']),
        "frecuencia_dias": cobro.get('frecuencia_dias'),
        "titulo": titulo_cobro(cobro),
    }


def estilo_tabla(color_encabezado, color_filas):
    return TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(color_encabezado)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor(color_filas)]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ])


def pdf_cobro(cobro):
    """PDF de estado de cuenta: resumen Deuda/Pagado/Saldo y tabla Pago # | Fecha | Monto | Metodo | Saldo Restante."""
    hoy = datetime.now()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=15,
                                 textColor=colors.HexColor('#1e40af'), spaceAfter=6)
    h2 = ParagraphStyle('H2', parent=styles['Heading2'], textColor=colors.HexColor('#1e40af'))
    story = [
        Paragraph(titulo_cobro(cobro), title_style),
        Paragraph(f"Generado: {hoy.strftime('%d/%m/%Y')}", styles['Normal']),
        Spacer(1, 0.25*inch),
    ]

    resumen = Table([
        ['Deuda', 'Pagado', 'Saldo'],
        [f"${cobro['deuda']:,.2f}", f"${cobro['pagado']:,.2f}", f"${cobro['saldo']:,.2f}"],
    ], colWidths=[2*inch, 2*inch, 2*inch])
    estilo = estilo_tabla('#1e40af', '#f3f4f6')
    estilo.add('FONTSIZE', (0, 1), (-1, 1), 13)
    estilo.add('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold')
    if cobro['saldo'] > 0:
        estilo.add('TEXTCOLOR', (2, 1), (2, 1), colors.HexColor('#dc2626'))
    resumen.setStyle(estilo)
    story.append(resumen)
    if cobro.get('frecuencia_dias'):
        story.append(Spacer(1, 0.1*inch))
        story.append(Paragraph(f"Paga cada {cobro['frecuencia_dias']} días promedio", styles['Normal']))
    story.append(Spacer(1, 0.3*inch))

    story.append(Paragraph(f"Historial de Pagos ({len(cobro['pagos'])})", h2))
    data = [['Pago #', 'Fecha', 'Monto', 'Metodo', 'Saldo Restante']]
    for idx, fila in enumerate(cobro['filas'], 1):
        data.append([str(idx), fila['fecha_corta'], f"${fila['monto']:,.2f}", fila['metodo'].capitalize(), f"${fila['saldo']:,.2f}"])
    data.append(['', 'TOTAL', f"${cobro['pagado']:,.2f}", '', f"${cobro['saldo']:,.2f}"])
    tabla = Table(data, colWidths=[0.7*inch, 1.1*inch, 1.3*inch, 1.6*inch, 1.5*inch])
    estilo = estilo_tabla('#059669', '#f0fdf4')
    for idx, fila in enumerate(cobro['filas'], 1):
        if fila['metodo'] == 'no especificado':
            estilo.add('TEXTCOLOR', (3, idx), (3, idx), colors.HexColor('#9ca3af'))
    estilo.add('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold')
    estilo.add('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e5e7eb'))
    tabla.setStyle(estilo)
    story.append(tabla)

    doc.build(story)
    return buffer.getvalue()


def excel_cobro(cobro):
    """Excel con 2 hojas: Pagos (Fecha, Monto, Metodo, Nota, Saldo) y Resumen."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    negrita = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1E40AF")
    wb = Workbook()

    def encabezado(ws, columnas):
        ws.append(columnas)
        for celda in ws[ws.max_row]:
            celda.font = negrita
            celda.fill = relleno

    ws = wb.active
    ws.title = "Pagos"
    encabezado(ws, ["Fecha", "Monto", "Metodo", "Nota", "Saldo"])
    gris = Font(color="9CA3AF", italic=True)
    for fila in cobro['filas']:
        fecha = datetime.strptime(fila['fecha'], '%Y-%m-%d') if fila['fecha'] else "sin fecha"
        ws.append([fecha, fila['monto'], fila['metodo'], fila['nota'], fila['saldo']])
        if fila['fecha']:
            ws.cell(row=ws.max_row, column=1).number_format = 'DD-MM-YY'
        if fila['metodo'] == 'no especificado':
            ws.cell(row=ws.max_row, column=3).font = gris
    ws.append(["TOTAL PAGADO", cobro['pagado'], "", "", cobro['saldo']])
    ws[ws.max_row][0].font = Font(bold=True)

    resumen = wb.create_sheet("Resumen")
    encabezado(resumen, ["Concepto", "Valor"])
    resumen.append(["Cliente", cobro['cliente']])
    resumen.append(["Deuda", cobro['deuda']])
    resumen.append(["Pagado", cobro['pagado']])
    resumen.append(["Saldo", cobro['saldo']])
    resumen.append(["Número de pagos", len(cobro['pagos'])])
    if cobro.get('frecuencia_dias'):
        resumen.append(["Frecuencia", f"Paga cada {cobro['frecuencia_dias']} días promedio"])

    for hoja in (ws, resumen):
        for col in hoja.columns:
            hoja.column_dimensions[col[0].column_letter].width = max(12, max(len(str(c.value or '')) for c in col) + 2)
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                if isinstance(celda.value, (int, float)) and celda.column_letter != 'A':
                    celda.number_format = '"$"#,##0.00'

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def registro_para_guardar(cobro, factura_path=""):
    """Lo que se guarda en memoria_global.json y en data/{tel}/cobro_deuda.json (campos deuda y pagado)."""
    return {
        "tipo": "cobro_deuda",
        "cliente": cobro['cliente'],
        "deuda": cobro['deuda'],
        "pagado": cobro['pagado'],
        "saldo": cobro['saldo'],
        "pagos": cobro['pagos'],
        "descripcion": cobro.get('descripcion') or '',
        "hash": cobro.get('hash', ''),
        "fecha": datetime.now().isoformat(),
        "factura_path": factura_path,
        "ultima_pregunta": "dashboard",
    }


def generar_reporte(datos_calculados, phone, guardar=None, factura_path="", incluir_archivos=True):
    """
    Entrada: el cobro que armó la Calculadora y el teléfono del usuario.
    guardar(phone, registro): función que lo escribe en memoria_global.json (no se guarda si es duplicado).
    Salida: {"mensaje_wa", "dashboard_cards", "pdf_bytes", "excel_bytes"}.
    """
    duplicado = bool(datos_calculados.get('duplicado'))
    if guardar and not duplicado:
        guardar(phone, registro_para_guardar(datos_calculados, factura_path))
    return {
        "mensaje_wa": mensaje_whatsapp(datos_calculados, duplicado),
        "dashboard_cards": dashboard_cards(datos_calculados),
        "pdf_bytes": pdf_cobro(datos_calculados) if incluir_archivos else None,
        "excel_bytes": excel_cobro(datos_calculados) if incluir_archivos else None,
    }
