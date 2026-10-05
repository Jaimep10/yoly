#!/usr/bin/env python
"""Contrato del reporte web editable: /reporte/{wa_id} y /api/gastos (app/web/reporte.py)."""
import io
import json
import os

import pytest
from openpyxl import load_workbook

import main
from app.services import reporte_service
from app.web import reporte

TEL = "whatsapp:+593991234567"
WA_ID = "593991234567"
OTRO = "593998887777"


def _gastos_iniciales():
    return [
        {"id": "g1", "fecha": "2026-10-01", "descripcion": "Supermaxi", "monto": 45.5, "categoria": "comida"},
        {"id": "g2", "fecha": "2026-10-02", "descripcion": "Renta", "monto": 380, "categoria": "vivienda"},
        {"fecha": "2026-10-03", "descripcion": "Taxi sin id", "monto": 12, "categoria": "transporte"},
    ]


@pytest.fixture
def web(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "DATA_DIR", str(tmp_path))
    monkeypatch.setenv("YOLY_WEB_SECRET", "secreto-de-prueba")
    # El bot guarda en data/whatsapp:+593.../principal; otro cliente con su propio gasto g1
    for carpeta, datos in ((TEL, _gastos_iniciales()),
                           (OTRO, [{"id": "g1", "fecha": "2026-10-01", "descripcion": "Del otro", "monto": 99, "categoria": "x"}])):
        ruta = tmp_path / carpeta / "principal"
        ruta.mkdir(parents=True)
        (ruta / "gastos.json").write_text(json.dumps(datos), encoding="utf-8")
    (tmp_path / TEL / "principal" / "ingresos.json").write_text(json.dumps([{"monto": 1000}]), encoding="utf-8")
    return main.app.test_client(), tmp_path


def _guardados(tmp_path, carpeta=TEL):
    return json.loads((tmp_path / carpeta / "principal" / "gastos.json").read_text(encoding="utf-8"))


def _t(wa_id=WA_ID):
    return reporte.token_para(wa_id)


def test_get_devuelve_gastos_y_totales(web):
    cliente, tmp_path = web
    r = cliente.get(f"/api/gastos/{WA_ID}?t={_t()}")
    assert r.status_code == 200
    datos = r.get_json()
    assert [g["descripcion"] for g in datos["gastos"]] == ["Supermaxi", "Renta", "Taxi sin id"]
    assert datos["totales"] == {"ingresos": 1000.0, "gastos": 437.5, "saldo": 562.5, "cantidad": 3}
    # El gasto sin id recibe uno y queda guardado, para poder editarlo/borrarlo
    assert all(g.get("id") for g in _guardados(tmp_path))


def test_borrar_desde_web_desaparece_del_almacenamiento(web):
    cliente, tmp_path = web
    r = cliente.delete(f"/api/gastos/g2?wa_id={WA_ID}&t={_t()}")
    assert r.status_code == 200
    assert r.get_json()["totales"]["gastos"] == 57.5
    assert "g2" not in [g["id"] for g in _guardados(tmp_path)]
    # Y ya no vuelve en el GET
    ids = [g["id"] for g in cliente.get(f"/api/gastos/{WA_ID}?t={_t()}").get_json()["gastos"]]
    assert "g2" not in ids and "g1" in ids


def test_editar_monto_y_categoria(web):
    cliente, tmp_path = web
    r = cliente.put(f"/api/gastos/g1?wa_id={WA_ID}&t={_t()}", json={"monto": "50.25", "categoria": "super"})
    assert r.status_code == 200
    assert r.get_json()["totales"]["gastos"] == 442.25
    g1 = next(g for g in _guardados(tmp_path) if g["id"] == "g1")
    assert g1["monto"] == 50.25 and g1["categoria"] == "super"
    assert g1["descripcion"] == "Supermaxi"  # lo demás no se toca


@pytest.mark.parametrize("monto", ["abc", "-5", "0", "nan", "inf"])
def test_editar_rechaza_montos_invalidos(web, monto):
    cliente, tmp_path = web
    r = cliente.put(f"/api/gastos/g1?wa_id={WA_ID}&t={_t()}", json={"monto": monto})
    assert r.status_code == 400
    assert next(g for g in _guardados(tmp_path) if g["id"] == "g1")["monto"] == 45.5


def test_cliente_no_borra_gastos_de_otro_wa_id(web):
    cliente, tmp_path = web
    # Con su propio token, el cliente intenta borrar g1 diciendo que es el otro número
    r = cliente.delete(f"/api/gastos/g1?wa_id={OTRO}&t={_t()}")
    assert r.status_code == 403
    # Con su token y su wa_id borra SU g1, nunca el del otro
    assert cliente.delete(f"/api/gastos/g1?wa_id={WA_ID}&t={_t()}").status_code == 200
    assert [g["id"] for g in _guardados(tmp_path, OTRO)] == ["g1"]
    # Un id que no es suyo no existe para él
    assert cliente.put(f"/api/gastos/g1?wa_id={WA_ID}&t={_t()}", json={"monto": 1}).status_code == 404


def test_sin_token_o_token_falso_no_entra(web):
    cliente, tmp_path = web
    assert cliente.get(f"/api/gastos/{WA_ID}").status_code == 403
    assert cliente.get(f"/api/gastos/{WA_ID}?t=falso").status_code == 403
    assert cliente.delete(f"/api/gastos/g1?wa_id={WA_ID}&t=falso").status_code == 403
    assert cliente.put(f"/api/gastos/g1?wa_id={WA_ID}", json={"monto": 1}).status_code == 403
    assert cliente.get(f"/reporte/{WA_ID}?t=falso").status_code == 403
    assert len(_guardados(tmp_path)) == 3


def test_sin_secreto_configurado_nadie_edita(web, monkeypatch):
    cliente, _ = web
    monkeypatch.delenv("YOLY_WEB_SECRET")
    monkeypatch.delenv("TWILIO_AUTH_TOKEN", raising=False)
    assert reporte.token_para(WA_ID) is None
    assert cliente.delete(f"/api/gastos/g1?wa_id={WA_ID}&t=").status_code == 403


def test_pagina_reporte_con_logo_y_tabla(web):
    cliente, _ = web
    r = cliente.get(f"/reporte/{WA_ID}?t={_t()}")
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    assert "YOLY ASISTENTE CONTABLE" in html
    for col in ("Fecha", "Descripción", "Monto", "Categoría"):
        assert col in html
    assert "fetch(" in html and "'DELETE'" in html and "'PUT'" in html


def test_link_del_bot_lleva_token_valido():
    os.environ["YOLY_WEB_SECRET"] = "secreto-de-prueba"
    try:
        link = reporte.link_reporte("https://yoly.test", TEL)
        assert link == f"https://yoly.test/reporte/{WA_ID[-10:]}?t={reporte.token_para(WA_ID)}"
        assert link in main.mensaje_link_dashboard("https://yoly.test", WA_ID)
        # El token sirve con cualquier forma del número (whatsapp:+593..., 593..., últimos 10)
        assert reporte.token_valido(WA_ID[-10:], reporte.token_para(TEL))
    finally:
        del os.environ["YOLY_WEB_SECRET"]


def test_excel_totales_son_formulas(web):
    cliente, _ = web
    gastos = [{"fecha": "2026-10-01", "descripcion": "A", "monto": 10, "categoria": "x"},
              {"fecha": "2026-10-02", "descripcion": "B", "monto": 5.5, "categoria": "y"}]
    wb = reporte_service.generar_excel(WA_ID, gastos, [{"fecha": "2026-10-01", "descripcion": "Sueldo", "monto": 100}])
    assert wb.sheetnames == ["RESUMEN", "GASTOS", "INGRESOS"]
    assert wb["GASTOS"]["C5"].value == "=SUM(C2:C4)"
    assert wb["INGRESOS"]["C4"].value == "=SUM(C2:C3)"
    assert wb["RESUMEN"]["B4"].value == "=INGRESOS!C4"
    assert wb["RESUMEN"]["B5"].value == "=GASTOS!C5"
    assert wb["RESUMEN"]["B6"].value == "=B4-B5"

    # Sin movimientos la fórmula no se apunta a sí misma ni al encabezado
    vacio = reporte_service.generar_excel(WA_ID, [], [])
    assert vacio["GASTOS"]["C3"].value == "=SUM(C2:C2)"

    # Descarga desde la página
    r = cliente.get(f"/reporte/{WA_ID}/excel?t={_t()}")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert load_workbook(io.BytesIO(r.data))["RESUMEN"]["B6"].value == "=B4-B5"
