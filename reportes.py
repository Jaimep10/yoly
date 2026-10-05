# Reportes por período y carpetas por tipo de transacción (sueldo, renta, deuda...).
# Lógica pura (sin Flask ni Twilio) para poder probarla sola, igual que precios.py.
import io
import json
import os
import re
import unicodedata
from calendar import monthrange
from datetime import date, datetime, timedelta

MESES = {
    'enero': 1, 'ene': 1, 'febrero': 2, 'feb': 2, 'marzo': 3, 'mar': 3, 'abril': 4, 'abr': 4,
    'mayo': 5, 'may': 5, 'junio': 6, 'jun': 6, 'julio': 7, 'jul': 7, 'agosto': 8, 'ago': 8,
    'septiembre': 9, 'setiembre': 9, 'sept': 9, 'sep': 9, 'octubre': 10, 'oct': 10,
    'noviembre': 11, 'nov': 11, 'diciembre': 12, 'dic': 12,
}
NOMBRE_MES = ['', 'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
              'septiembre', 'octubre', 'noviembre', 'diciembre']
MES_CORTO = ['', 'ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov', 'dic']

# Nombres largos primero para que "junio" no se lea como "jun" + "io"
_MES_RE = '(' + '|'.join(sorted(MESES, key=len, reverse=True)) + r')\b'
_AL = r'\s*(?:al|a|hasta|-|–)\s*'
_ANIO = r'(?:\s*(?:de|del)?\s*(\d{4}))?'


def hoy_fecha():
    return date.today()


def sin_acentos(texto):
    return ''.join(c for c in unicodedata.normalize('NFD', str(texto or '')) if unicodedata.category(c) != 'Mn')


def _normalizar(texto):
    return re.sub(r'\s+', ' ', sin_acentos(texto).lower()).strip()


def _dia(anio, mes, dia):
    """Fecha con el día ajustado al último del mes ("31 de junio" -> 30 de junio)."""
    return date(anio, mes, min(max(dia, 1), monthrange(anio, mes)[1]))


def _anio_para(mes, anio_texto, hoy):
    """Sin año escrito: el mes más reciente que ya pasó o está en curso."""
    if anio_texto:
        return int(anio_texto)
    return hoy.year if mes <= hoy.month else hoy.year - 1


def etiqueta_periodo(inicio, fin):
    """'1 al 20 de junio 2026', 'junio 2026', '15 may al 10 jun 2026'"""
    if inicio.day == 1 and fin == _dia(fin.year, fin.month, 31) and (inicio.year, inicio.month) == (fin.year, fin.month):
        return f"{NOMBRE_MES[inicio.month]} {inicio.year}"
    if (inicio.year, inicio.month) == (fin.year, fin.month):
        if inicio == fin:
            return f"{inicio.day} de {NOMBRE_MES[inicio.month]} {inicio.year}"
        return f"{inicio.day} al {fin.day} de {NOMBRE_MES[inicio.month]} {inicio.year}"
    if inicio.year == fin.year:
        return f"{inicio.day} {MES_CORTO[inicio.month]} al {fin.day} {MES_CORTO[fin.month]} {fin.year}"
    return f"{inicio.day} {MES_CORTO[inicio.month]} {inicio.year} al {fin.day} {MES_CORTO[fin.month]} {fin.year}"


def _periodo(inicio, fin):
    if fin < inicio:
        inicio, fin = fin, inicio
    return {"inicio": inicio, "fin": fin, "etiqueta": etiqueta_periodo(inicio, fin)}


def parsear_periodo(texto, hoy=None):
    """
    Entiende el período de un mensaje. Devuelve {"inicio", "fin", "etiqueta"} o None.
    "del 1 al 20 de junio", "junio 1-20", "15 de mayo al 10 de junio", "1/6 al 20/6",
    "junio", "junio 2025", "este mes", "mes pasado", "esta semana", "semana pasada",
    "ultimos 15 dias", "primera quincena de junio", "hoy", "ayer", "este año".
    Fechas con barra se leen día/mes (1/6 = 1 de junio).
    """
    hoy = hoy or hoy_fecha()
    if isinstance(hoy, datetime):
        hoy = hoy.date()
    t = _normalizar(texto)

    # "del 15 de mayo al 10 de junio [2026]"
    m = re.search(r'(\d{1,2})\s*(?:de\s+)?' + _MES_RE + _ANIO + _AL + r'(\d{1,2})\s*(?:de\s+)?' + _MES_RE + _ANIO, t)
    if m:
        d1, mes1, a1, d2, mes2, a2 = m.groups()
        mes1, mes2 = MESES[mes1], MESES[mes2]
        anio2 = _anio_para(mes2, a2, hoy)
        anio1 = int(a1) if a1 else (anio2 if mes1 <= mes2 else anio2 - 1)
        return _periodo(_dia(anio1, mes1, int(d1)), _dia(anio2, mes2, int(d2)))

    # "del 1 al 20 de junio [2026]", "1-20 junio"
    m = re.search(r'(\d{1,2})' + _AL + r'(\d{1,2})\s*(?:de\s+)?' + _MES_RE + _ANIO, t)
    if m:
        d1, d2, mes, a = m.groups()
        mes = MESES[mes]
        anio = _anio_para(mes, a, hoy)
        return _periodo(_dia(anio, mes, int(d1)), _dia(anio, mes, int(d2)))

    # "junio 1-20", "junio del 1 al 20", "junio 2026 del 1 al 20"
    m = re.search(_MES_RE + r'\s*(?:de\s+)?(?:(\d{4})\s*)?(?:del?\s+)?(\d{1,2})' + _AL + r'(\d{1,2})\b', t)
    if m:
        mes, a, d1, d2 = m.groups()
        mes = MESES[mes]
        anio = _anio_para(mes, a, hoy)
        return _periodo(_dia(anio, mes, int(d1)), _dia(anio, mes, int(d2)))

    # "1/6 al 20/6", "01-06-2026 al 20-06-2026" (día/mes)
    m = re.search(r'(\d{1,2})[/.](\d{1,2})(?:[/.](\d{2,4}))?' + _AL + r'(\d{1,2})[/.](\d{1,2})(?:[/.](\d{2,4}))?', t)
    if m:
        d1, m1, a1, d2, m2, a2 = m.groups()
        try:
            m1, m2 = int(m1), int(m2)
            if not (1 <= m1 <= 12 and 1 <= m2 <= 12):
                raise ValueError
            def anio(a, mes):
                if not a:
                    return _anio_para(mes, None, hoy)
                return int(a) + 2000 if len(a) == 2 else int(a)
            fin = _dia(anio(a2, m2), m2, int(d2))
            inicio_anio = anio(a1, m1) if a1 else (fin.year if m1 <= m2 else fin.year - 1)
            return _periodo(_dia(inicio_anio, m1, int(d1)), fin)
        except ValueError:
            pass

    # Quincenas: "primera quincena de junio", "segunda quincena", "esta quincena"
    m = re.search(r'(primera|1ra|segunda|2da|esta)\s+quincena(?:\s+(?:de|del)\s+' + _MES_RE + _ANIO + ')?', t)
    if m:
        cual, mes, a = m.groups()
        if mes:
            mes = MESES[mes]
            anio = _anio_para(mes, a, hoy)
        else:
            mes, anio = hoy.month, hoy.year
            if 'pasad' in t:
                mes, anio = (12, anio - 1) if mes == 1 else (mes - 1, anio)
        if cual == 'esta':
            cual = 'primera' if hoy.day <= 15 else 'segunda'
        if cual in ('primera', '1ra'):
            return _periodo(date(anio, mes, 1), date(anio, mes, 15))
        return _periodo(date(anio, mes, 16), _dia(anio, mes, 31))

    m = re.search(r'ultim[oa]s?\s+(\d{1,3})\s+dias', t)
    if m:
        return _periodo(hoy - timedelta(days=int(m.group(1)) - 1), hoy)

    lunes = hoy - timedelta(days=hoy.weekday())
    if re.search(r'\bsemana pasada\b', t):
        return _periodo(lunes - timedelta(days=7), lunes - timedelta(days=1))
    if re.search(r'\besta semana\b', t):
        return _periodo(lunes, lunes + timedelta(days=6))
    if re.search(r'\bmes pasado\b', t):
        fin = date(hoy.year, hoy.month, 1) - timedelta(days=1)
        return _periodo(date(fin.year, fin.month, 1), fin)
    if re.search(r'\beste mes\b', t):
        return _periodo(date(hoy.year, hoy.month, 1), _dia(hoy.year, hoy.month, 31))
    if re.search(r'\beste ano\b', t):
        return _periodo(date(hoy.year, 1, 1), date(hoy.year, 12, 31))
    if re.search(r'\bayer\b', t):
        return _periodo(hoy - timedelta(days=1), hoy - timedelta(days=1))
    if re.search(r'\bhoy\b', t):
        return _periodo(hoy, hoy)

    # Mes solo: "junio", "junio 2025", "mes de junio" (abreviaturas solo con año para no confundir "mar")
    m = re.search(r'\b(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|setiembre|octubre|noviembre|diciembre)\b'
                  + _ANIO, t)
    if m:
        mes, a = MESES[m.group(1)], m.group(2)
        anio = _anio_para(mes, a, hoy)
        return _periodo(date(anio, mes, 1), _dia(anio, mes, 31))

    # "del 1 al 20" sin mes: este mes
    m = re.search(r'\bdel?\s+(\d{1,2})\s+al?\s+(\d{1,2})\b', t)
    if m:
        return _periodo(_dia(hoy.year, hoy.month, int(m.group(1))), _dia(hoy.year, hoy.month, int(m.group(2))))
    return None


# ==================== CLASIFICACIÓN DE TRANSFERENCIAS ====================

# Carpeta -> palabras del concepto. El orden decide si hay más de una.
CARPETAS = [
    ('sueldo', ['sueldo', 'sueldos', 'salario', 'salarios', 'nomina', 'quincena', 'jornal', 'pago empleado',
                'pago empleada', 'pago trabajador', 'honorarios', 'paga semanal', 'payroll', 'salary', 'wages']),
    ('renta', ['renta', 'arriendo', 'alquiler', 'arrendamiento', 'rent', 'casero', 'casera']),
    ('deuda', ['deuda', 'prestamo', 'prestamos', 'abono', 'cuota', 'credito', 'loan', 'debo', 'adeudo',
               'tarjeta de credito']),
    ('servicios', ['luz', 'agua', 'internet', 'telefono', 'celular', 'electricidad', 'cable', 'gas natural',
                   'plan movil']),
]
CARPETA_INGRESO = 'ingreso'
CARPETA_REVISAR = 'por_revisar'
CARPETA_COMPRAS = 'compras'
TODAS_CARPETAS = [c for c, _ in CARPETAS] + [CARPETA_INGRESO, CARPETA_COMPRAS, CARPETA_REVISAR]
EMOJI_CARPETA = {'sueldo': '👷', 'renta': '🏠', 'deuda': '💳', 'servicios': '💡', 'ingreso': '💰',
                 'compras': '🛒', 'por_revisar': '❓'}


def clasificar_concepto(texto):
    """Carpeta según las palabras del concepto ('Sueldo mensual Juan' -> 'sueldo'); None si no hay pista."""
    t = _normalizar(texto)
    if not t:
        return None
    for carpeta, palabras in CARPETAS:
        if any(re.search(r'\b' + re.escape(p) + r'\b', t) for p in palabras):
            return carpeta
    return None


def clasificar_transferencia(concepto, direccion=None, texto_usuario=''):
    """
    Decide carpeta y si es ingreso o gasto.
    Lo que escribe el usuario junto a la foto manda sobre el concepto del banco.
    Recibida sin pista -> 'ingreso'. Enviada sin pista -> 'por_revisar' (no adivinamos).
    Si no se sabe si entró o salió (ni la foto ni el usuario lo dicen) -> 'por_revisar' con
    direccion_desconocida=True: no cuenta ni como ingreso ni como gasto hasta que el usuario diga.
    """
    carpeta = clasificar_concepto(texto_usuario) or clasificar_concepto(concepto)
    direccion = direccion_por_texto(texto_usuario) or str(direccion or '').lower()
    recibida = direccion.startswith('recib')
    if not direccion.startswith(('recib', 'envi')):
        return {"carpeta": CARPETA_REVISAR, "movimiento": 'gasto', "direccion_desconocida": True}
    if not carpeta:
        carpeta = CARPETA_INGRESO if recibida else CARPETA_REVISAR
    return {"carpeta": carpeta, "movimiento": 'ingreso' if recibida else 'gasto', "direccion_desconocida": False}


_RE_ENTRA = re.compile(r'\b(recibi|recibido|recibida|me pagaron|me pago|me depositaron|me deposito|me transfirieron|'
                       r'me transfirio|me enviaron|me envio|me mandaron|me mando|cobre|cobrado|vendi|venta|ventas|'
                       r'ingreso|ingresos|entrada|gane|abonaron)\b')
_RE_SALE = re.compile(r'\b(pague|pagado|gaste|compre|envie|mande|transferi|deposite|egreso|salida|gasto|gastos|compra)\b')


def direccion_por_texto(texto):
    """'recibida' / 'enviada' según lo que escribió el usuario; '' si no lo dice o dice las dos cosas."""
    t = _normalizar(texto)
    entra, sale = bool(_RE_ENTRA.search(t)), bool(_RE_SALE.search(t))
    if entra == sale:
        return ''
    return 'recibida' if entra else 'enviada'


def ruta_carpeta(data_dir, telefono, carpeta, cuenta="principal"):
    ruta = os.path.join(data_dir, telefono, cuenta, carpeta)
    os.makedirs(ruta, exist_ok=True)
    return os.path.join(ruta, 'transacciones.json')


def _leer(archivo):
    if not os.path.exists(archivo):
        return []
    try:
        with open(archivo, 'r', encoding='utf-8') as f:
            datos = json.load(f)
        return datos if isinstance(datos, list) else []
    except Exception:
        return []


def _escribir(archivo, datos):
    os.makedirs(os.path.dirname(archivo), exist_ok=True)
    with open(archivo, 'w', encoding='utf-8') as f:
        json.dump(datos, f, ensure_ascii=False, indent=2)


def guardar_en_carpeta(data_dir, telefono, carpeta, transaccion, cuenta="principal"):
    """Agrega la transacción a data/{telefono}/{cuenta}/{carpeta}/transacciones.json"""
    archivo = ruta_carpeta(data_dir, telefono, carpeta, cuenta)
    datos = _leer(archivo)
    datos.append(dict(transaccion, carpeta=carpeta))
    _escribir(archivo, datos)
    return archivo


def _carpetas_usuario(data_dir, telefono, cuenta="principal"):
    """Carpetas de datos del usuario por últimos 10 dígitos, dentro de la cuenta especificada."""
    ultimos10 = re.sub(r'[^0-9]', '', telefono or '')[-10:]
    if not ultimos10 or not os.path.isdir(data_dir):
        return []
    # Buscar data_dir/{phone}/{cuenta}/
    resultado = []
    for c in sorted(os.listdir(data_dir)):
        if re.sub(r'[^0-9]', '', c)[-10:] == ultimos10:
            cuenta_path = os.path.join(data_dir, c, cuenta)
            if os.path.isdir(cuenta_path):
                resultado.append(cuenta_path)
    return resultado


def cargar_carpetas(data_dir, telefono, cuenta="principal"):
    """{carpeta: [transacciones]} solo con las carpetas que tienen algo."""
    resultado = {}
    for base in _carpetas_usuario(data_dir, telefono, cuenta):
        for carpeta in TODAS_CARPETAS:
            datos = _leer(os.path.join(base, carpeta, 'transacciones.json'))
            if datos:
                resultado.setdefault(carpeta, []).extend(datos)
    for carpeta in resultado:
        resultado[carpeta].sort(key=lambda x: x.get('fecha') or '')
    return resultado


def cargar_movimientos(data_dir, telefono, cuenta="principal"):
    """(gastos, ingresos) del usuario, de todas sus carpetas (gastos.json e ingresos.json)."""
    gastos, ingresos = [], []
    for base in _carpetas_usuario(data_dir, telefono, cuenta):
        gastos.extend(_leer(os.path.join(base, 'gastos.json')))
        ingresos.extend(_leer(os.path.join(base, 'ingresos.json')))
    return gastos, ingresos


def ultima_transferencia(data_dir, telefono, cuenta="principal"):
    """(base, carpeta, transacción) de la última transferencia guardada; prefiere las por revisar."""
    candidatos = []
    for base in _carpetas_usuario(data_dir, telefono, cuenta):
        for carpeta in TODAS_CARPETAS:
            if carpeta == CARPETA_COMPRAS:
                continue
            for t in _leer(os.path.join(base, carpeta, 'transacciones.json')):
                candidatos.append((carpeta == CARPETA_REVISAR, t.get('timestamp') or '', base, carpeta, t))
    if not candidatos:
        return None
    _, _, base, carpeta, t = max(candidatos, key=lambda c: (c[0], c[1]))
    return base, carpeta, t


def mover_transaccion(data_dir, telefono, nueva_carpeta, cuenta="principal"):
    """
    Mueve la última transferencia (la por revisar primero) a otra carpeta y actualiza su
    categoría en gastos.json / ingresos.json. Devuelve (transacción, carpeta_anterior) o None.
    """
    encontrada = ultima_transferencia(data_dir, telefono, cuenta)
    if not encontrada:
        return None
    base, carpeta, t = encontrada
    origen = os.path.join(base, carpeta, 'transacciones.json')
    _escribir(origen, [x for x in _leer(origen) if x.get('id') != t.get('id')])
    if nueva_carpeta == CARPETA_INGRESO:
        t['movimiento'] = 'ingreso'
    # El usuario ya dijo qué es: deja de estar en duda
    t['direccion_desconocida'] = False
    destino = os.path.join(base, nueva_carpeta, 'transacciones.json')
    _escribir(destino, _leer(destino) + [dict(t, carpeta=nueva_carpeta, revisar=False)])

    # La lista de gastos/ingresos (de donde salen los reportes) se mantiene igual a la carpeta
    for base_lista in _carpetas_usuario(data_dir, telefono, cuenta):
        for nombre in ('gastos.json', 'ingresos.json'):
            archivo = os.path.join(base_lista, nombre)
            lista = _leer(archivo)
            movido = next((x for x in lista if x.get('id') == t.get('id')), None)
            if not movido:
                continue
            restantes = [x for x in lista if x.get('id') != t.get('id')]
            movido = dict(movido, categoria=nueva_carpeta, movimiento=t.get('movimiento'), revisar=False,
                          direccion_desconocida=False)
            nombre_destino = 'ingresos.json' if t.get('movimiento') == 'ingreso' else 'gastos.json'
            if nombre_destino == nombre:
                _escribir(archivo, restantes + [movido])
            else:
                _escribir(archivo, restantes)
                otro = os.path.join(base_lista, nombre_destino)
                _escribir(otro, _leer(otro) + [movido])
    return t, carpeta


_RE_MOVER = re.compile(
    r'^(?:no,?\s+)?(?:es|era|fue|eso es|esa es|ese es|esto es|ponlo|ponla|pon|muevelo|muevela|mueve|guardalo|guardala|'
    r'mandalo|mandala|va|son)\s+(?:a|en|de|para)?\s*(?:la\s+carpeta\s+(?:de\s+)?)?(?:(?:un|una|el|la|mi)\s+)?(?:pago\s+(?:de\s+)?(?:la\s+|el\s+|mi\s+)?)?'
    r'(sueldo|salario|renta|arriendo|alquiler|deuda|prestamo|servicios?|ingreso|luz|agua|internet)\b')


def pedido_mover(texto):
    """'es renta', 'ponlo en sueldo', 'era deuda' -> carpeta; None si el mensaje no es eso."""
    t = _normalizar(texto).rstrip('.!')
    if len(t.split()) > 8:
        return None
    m = _RE_MOVER.match(t)
    if not m:
        return None
    palabra = m.group(1)
    if palabra in ('ingreso',):
        return CARPETA_INGRESO
    if palabra.startswith('servicio'):
        return 'servicios'
    return clasificar_concepto(palabra)


# ==================== MOVIMIENTO ESCRITO ("me pagaron 300", "gasté 45.50 en comida") ====================

# Solo verbos en pasado: "gano 1200 al mes" es un presupuesto, no un movimiento de hoy.
_RE_VERBO_ENTRA = re.compile(r'\b(recibi|me pagaron|me pago|me depositaron|me deposito|me transfirieron|me transfirio|'
                             r'me enviaron|me envio|me mandaron|me mando|cobre|vendi|gane|me entraron|entraron)\b')
_RE_VERBO_SALE = re.compile(r'\b(pague|gaste|compre|envie|mande|transferi|deposite)\b')
_RE_FECHA_NUM = re.compile(r'\b\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?\b')
_RE_FECHA_MES = re.compile(r'\b(\d{1,2})\s+de\s+' + _MES_RE + r'(?:\s+(?:de|del)\s+(\d{4}))?')
_RE_MONTO = re.compile(r'(\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?|\d+(?:[.,]\d{1,2})?)')


def _monto_texto(numero):
    """'1.200' -> 1200, '1,200.50' -> 1200.5, '45,50' -> 45.5, '45.50' -> 45.5"""
    partes = re.split(r'[.,]', numero)
    if len(partes) > 1 and len(partes[-1]) in (1, 2):
        return float(''.join(partes[:-1]) + '.' + partes[-1])
    return float(''.join(partes))


def menciona_movimiento(texto):
    """True si el mensaje dice que entró o salió plata ("pagué...", "me pagaron..."), aunque no esté claro cuánto."""
    t = _normalizar(texto)
    return bool(_RE_VERBO_ENTRA.search(t) or _RE_VERBO_SALE.search(t))


def movimiento_de_texto(texto, hoy=None):
    """
    Un ingreso o un gasto escrito en una sola frase, sin IA:
    "me pagaron 300 por una venta" -> {"tipo": "ingreso", "monto": 300.0, ...}
    "gasté 45.50 en comida ayer"   -> {"tipo": "gasto", "monto": 45.5, ...}
    Devuelve None si no está claro (sin verbo, verbos de entrada y salida juntos, ningún monto o
    varios montos distintos): en ese caso no se guarda nada en vez de adivinar.
    """
    t = _normalizar(texto)
    entra, sale = bool(_RE_VERBO_ENTRA.search(t)), bool(_RE_VERBO_SALE.search(t))
    if entra == sale:
        return None
    hoy = hoy or hoy_fecha()

    fecha = hoy
    m = _RE_FECHA_MES.search(t)
    if m:
        anio = int(m.group(3)) if m.group(3) else hoy.year
        try:
            fecha = date(anio, MESES[m.group(2)], int(m.group(1)))
        except ValueError:
            return None
        t_sin_fecha = t[:m.start()] + ' ' + t[m.end():]
    else:
        t_sin_fecha = t
        f = _RE_FECHA_NUM.search(t)
        if f:
            partes = [int(x) for x in re.split(r'[/-]', f.group(0))]
            anio = partes[2] if len(partes) == 3 else hoy.year
            anio = anio + 2000 if anio < 100 else anio
            try:
                fecha = date(anio, partes[1], partes[0])
            except ValueError:
                return None
            t_sin_fecha = t[:f.start()] + ' ' + t[f.end():]
        elif re.search(r'\bayer\b', t):
            fecha = hoy - timedelta(days=1)
        elif re.search(r'\banteayer\b', t):
            fecha = hoy - timedelta(days=2)

    montos = {_monto_texto(n) for n in _RE_MONTO.findall(t_sin_fecha)}
    montos.discard(0.0)
    if len(montos) != 1:
        return None
    monto = montos.pop()

    tipo = 'ingreso' if entra else 'gasto'
    categoria = clasificar_concepto(t)
    if not categoria:
        if tipo == 'ingreso':
            categoria = 'ventas' if re.search(r'\b(vendi|venta|ventas)\b', t) else CARPETA_INGRESO
        else:
            categoria = 'otro'
    return {"tipo": tipo, "monto": round(monto, 2), "fecha": fecha.isoformat(), "categoria": categoria,
            "descripcion": re.sub(r'\s+', ' ', str(texto)).strip()[:80]}


# ==================== RESUMEN DEL PERÍODO ====================

def _a_numero(valor):
    if isinstance(valor, bool):
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    try:
        return float(re.sub(r'[^0-9.\-]', '', str(valor).replace(',', '')))
    except ValueError:
        return 0.0


def _fecha(texto):
    try:
        return datetime.strptime(str(texto)[:10], '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None


def _es_registro_pagos(g):
    """Fotos viejas de libreta de pagos guardadas como un gasto: no son gastos del período."""
    return isinstance(g.get('pagos'), list) and (len(g['pagos']) > 1 or g.get('cliente'))


def movimientos_periodo(gastos, ingresos, inicio, fin):
    """
    Lista plana de movimientos entre inicio y fin (incluidos):
    {fecha, tipo: gasto|ingreso, categoria, descripcion, monto}. Los desgloses
    ("envié 500, renta 380...") se abren en sus categorías. También cuenta los que no tienen fecha.
    """
    movs, sin_fecha = [], 0
    for g in gastos:
        if _es_registro_pagos(g):
            continue
        f = _fecha(g.get('fecha'))
        if not f:
            sin_fecha += 1
            continue
        if not (inicio <= f <= fin):
            continue
        desc = g.get('descripcion') or g.get('cliente') or 'Gasto'
        if isinstance(g.get('desglose_json'), dict) and g['desglose_json']:
            for cat, monto in g['desglose_json'].items():
                movs.append({"fecha": f, "tipo": "gasto", "categoria": str(cat).lower(),
                             "descripcion": desc, "monto": _a_numero(monto), "revisar": False})
            reserva = _a_numero(g.get('reserva'))
            if reserva > 0:
                movs.append({"fecha": f, "tipo": "gasto", "categoria": "reserva", "descripcion": desc,
                             "monto": reserva, "revisar": False})
            continue
        movs.append({"fecha": f, "tipo": "por_revisar" if g.get('direccion_desconocida') else "gasto",
                     "categoria": (g.get('categoria') or 'otro').lower(),
                     "descripcion": desc, "monto": _a_numero(g.get('monto')), "revisar": bool(g.get('revisar')),
                     "aproximada": bool(g.get('fecha_aproximada'))})
    for i in ingresos:
        f = _fecha(i.get('fecha'))
        if not f:
            sin_fecha += 1
            continue
        if inicio <= f <= fin:
            movs.append({"fecha": f, "tipo": "ingreso", "categoria": (i.get('categoria') or 'ingreso').lower(),
                         "descripcion": i.get('descripcion') or i.get('cliente') or 'Ingreso',
                         "monto": _a_numero(i.get('monto')), "revisar": bool(i.get('revisar')),
                         "aproximada": bool(i.get('fecha_aproximada'))})
    movs.sort(key=lambda m: (m['fecha'], m['tipo']))
    return movs, sin_fecha


def resumir(movs):
    """Totales con Python (nunca con la IA): ingresos, gastos, balance, por categoría y por semana."""
    gastos = [m for m in movs if m['tipo'] == 'gasto']
    ingresos = [m for m in movs if m['tipo'] == 'ingreso']
    total_g = round(sum(m['monto'] for m in gastos), 2)
    total_i = round(sum(m['monto'] for m in ingresos), 2)

    def por_cat(lista):
        acc = {}
        for m in lista:
            acc[m['categoria']] = acc.get(m['categoria'], 0) + m['monto']
        return sorted(((c, round(v, 2)) for c, v in acc.items()), key=lambda x: -x[1])

    semanas = {}
    for m in gastos + ingresos:
        lunes = m['fecha'] - timedelta(days=m['fecha'].weekday())
        s = semanas.setdefault(lunes, {"lunes": lunes, "domingo": lunes + timedelta(days=6), "ingresos": 0.0, "gastos": 0.0})
        s['ingresos' if m['tipo'] == 'ingreso' else 'gastos'] += m['monto']
    semanas = [dict(s, ingresos=round(s['ingresos'], 2), gastos=round(s['gastos'], 2),
                    balance=round(s['ingresos'] - s['gastos'], 2)) for _, s in sorted(semanas.items())]
    return {
        "total_ingresos": total_i, "total_gastos": total_g, "balance": round(total_i - total_g, 2),
        "gastos_por_categoria": por_cat(gastos), "ingresos_por_categoria": por_cat(ingresos),
        "semanas": semanas, "movimientos": movs, "num_gastos": len(gastos), "num_ingresos": len(ingresos),
        "por_revisar": sum(1 for m in movs if m.get('revisar') or m['categoria'] == CARPETA_REVISAR),
        # Sin saber si entró o salió: fuera de los totales hasta que el usuario lo diga
        "sin_direccion": len([m for m in movs if m['tipo'] == 'por_revisar']),
        "monto_sin_direccion": round(sum(m['monto'] for m in movs if m['tipo'] == 'por_revisar'), 2),
    }


def dinero(valor):
    return f"-${abs(valor):,.2f}" if valor < 0 else f"${valor:,.2f}"


def barra_texto(valor, maximo, ancho=8):
    llenos = int(round(ancho * valor / maximo)) if maximo else 0
    return '▓' * llenos + '░' * (ancho - llenos)


def texto_resumen(resumen, periodo, sin_fecha=0):
    """Mensaje de WhatsApp: totales, gráfico de barras en texto y desglose por categoría."""
    if not resumen['movimientos']:
        texto = f"📊 Del {periodo['etiqueta']} no tengo gastos ni ingresos guardados."
        if sin_fecha:
            texto += f"\n({sin_fecha} registros no tienen fecha y no los pude ubicar.)"
        return texto
    signo = '+' if resumen['balance'] >= 0 else '-'
    lineas = [
        f"📊 *Resumen del {periodo['etiqueta']}*",
        f"💰 Ingresos: {dinero(resumen['total_ingresos'])}",
        f"💸 Gastos: {dinero(resumen['total_gastos'])}",
        f"⚖️ Balance: {signo}{dinero(abs(resumen['balance']))}",
    ]
    cats = resumen['gastos_por_categoria']
    if cats:
        lineas += ["", "*Gastos por categoría:*"]
        maximo = cats[0][1]
        for cat, monto in cats[:8]:
            pct = 100 * monto / resumen['total_gastos'] if resumen['total_gastos'] else 0
            lineas.append(f"{barra_texto(monto, maximo)} {cat.replace('_', ' ')}: {dinero(monto)} ({pct:.0f}%)")
        if len(cats) > 8:
            lineas.append(f"… y {len(cats) - 8} categorías más")
    if len(resumen['semanas']) > 1:
        lineas += ["", "*Por semana (lun–dom):*"]
        for s in resumen['semanas']:
            lineas.append(f"{s['lunes'].strftime('%d/%m')}–{s['domingo'].strftime('%d/%m')}: "
                          f"+{dinero(s['ingresos'])} / -{dinero(s['gastos'])}")
    lineas.append("")
    lineas.append(f"{resumen['num_gastos']} gastos y {resumen['num_ingresos']} ingresos en el período.")
    if resumen['por_revisar']:
        lineas.append(f"❓ {resumen['por_revisar']} sin clasificar (marcados para revisar).")
    if resumen.get('sin_direccion'):
        lineas.append(f"❓ {resumen['sin_direccion']} movimiento(s) por {dinero(resumen['monto_sin_direccion'])} "
                      "no entraron en los totales: no sé si fueron ingreso o gasto.")
    if sin_fecha:
        lineas.append(f"⚠️ {sin_fecha} registros sin fecha no entraron.")
    return "\n".join(lineas)


# ==================== EXCEL Y PDF DEL PERÍODO ====================

def generar_excel(resumen, periodo, ruta):
    """Excel con hojas Resumen (con gráfico), Movimientos y Por semana. Valores fijos, sin fórmulas."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill

    negrita = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1E40AF")
    wb = Workbook()

    def encabezado(ws, columnas):
        ws.append(columnas)
        for celda in ws[ws.max_row]:
            celda.font, celda.fill = negrita, relleno

    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Resumen del {periodo['etiqueta']}"])
    ws['A1'].font = Font(bold=True, size=14, color="1E40AF")
    ws.append([])
    encabezado(ws, ["Concepto", "Monto"])
    ws.append(["Ingresos", resumen['total_ingresos']])
    ws.append(["Gastos", resumen['total_gastos']])
    ws.append(["Balance", resumen['balance']])
    ws.append([])
    encabezado(ws, ["Categoría de gasto", "Monto", "%"])
    fila_cat = ws.max_row + 1
    for cat, monto in resumen['gastos_por_categoria']:
        pct = round(100 * monto / resumen['total_gastos'], 1) if resumen['total_gastos'] else 0
        ws.append([cat, monto, pct])
    if resumen['gastos_por_categoria']:
        grafico = BarChart()
        grafico.type = "bar"
        grafico.title = "Gastos por categoría"
        grafico.legend = None
        grafico.add_data(Reference(ws, min_col=2, min_row=fila_cat - 1, max_row=ws.max_row), titles_from_data=True)
        grafico.set_categories(Reference(ws, min_col=1, min_row=fila_cat, max_row=ws.max_row))
        grafico.height, grafico.width = 8, 14
        ws.add_chart(grafico, "E3")
    if resumen['ingresos_por_categoria']:
        ws.append([])
        encabezado(ws, ["Categoría de ingreso", "Monto"])
        for cat, monto in resumen['ingresos_por_categoria']:
            ws.append([cat, monto])

    mv = wb.create_sheet("Movimientos")
    encabezado(mv, ["Fecha", "Tipo", "Categoría", "Descripción", "Monto", "Revisar"])
    for m in resumen['movimientos']:
        mv.append([m['fecha'], m['tipo'], m['categoria'], m['descripcion'], m['monto'],
                   "Revisar" if m.get('revisar') or m['categoria'] == CARPETA_REVISAR else ""])
        mv.cell(row=mv.max_row, column=1).number_format = 'DD/MM/YYYY'

    se = wb.create_sheet("Por semana")
    encabezado(se, ["Semana (lunes)", "Hasta (domingo)", "Ingresos", "Gastos", "Balance"])
    for s in resumen['semanas']:
        se.append([s['lunes'], s['domingo'], s['ingresos'], s['gastos'], s['balance']])
        for col in (1, 2):
            se.cell(row=se.max_row, column=col).number_format = 'DD/MM/YYYY'

    columnas_dinero = {ws: 'B', mv: 'E', se: 'CDE'}
    for hoja, letras in columnas_dinero.items():
        for col in hoja.columns:
            letra = col[0].column_letter
            hoja.column_dimensions[letra].width = max(12, min(45, max(len(str(c.value or '')) for c in col) + 2))
            if letra in letras:
                for celda in col[1:]:
                    if isinstance(celda.value, (int, float)):
                        celda.number_format = '"$"#,##0.00'
    ws.column_dimensions['A'].width = 24
    wb.save(ruta)
    return ruta


def grafico_png(resumen):
    """Barras horizontales de gastos por categoría (PNG en memoria) o None si no hay gastos."""
    cats = resumen['gastos_por_categoria'][:10]
    if not cats:
        return None
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    nombres = [c.replace('_', ' ') for c, _ in cats][::-1]
    montos = [m for _, m in cats][::-1]
    fig, ax = plt.subplots(figsize=(6.5, max(1.8, 0.45 * len(cats) + 0.8)), dpi=150)
    barras = ax.barh(nombres, montos, color="#3b5bdb", height=0.6)
    for barra, monto in zip(barras, montos):
        ax.text(barra.get_width(), barra.get_y() + barra.get_height() / 2, f"  {dinero(monto)}",
                va='center', fontsize=8, color="#1f2937")
    ax.set_xlim(0, max(montos) * 1.3)
    ax.xaxis.set_visible(False)
    for lado in ('top', 'right', 'bottom'):
        ax.spines[lado].set_visible(False)
    ax.tick_params(axis='y', labelsize=9, length=0)
    ax.set_title("Gastos por categoría", fontsize=11, loc='left', color="#1e40af")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png')
    plt.close(fig)
    buf.seek(0)
    return buf


def generar_pdf(resumen, periodo, ruta):
    """PDF con totales, gráfico por categoría, semanas y lista de movimientos."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    titulo = ParagraphStyle('T', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#1e40af'))
    h2 = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=12, textColor=colors.HexColor('#1e40af'))

    def estilo(color):
        return TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(color)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
            ('ALIGN', (-1, 1), (-1, -1), 'RIGHT'),
        ])

    story = [Paragraph(f"Resumen del {periodo['etiqueta']}", titulo),
             Paragraph(f"Generado el {date.today().strftime('%d/%m/%Y')}", styles['Normal']), Spacer(1, 0.2 * inch)]
    totales = Table([['Ingresos', 'Gastos', 'Balance'],
                     [dinero(resumen['total_ingresos']), dinero(resumen['total_gastos']), dinero(resumen['balance'])]],
                    colWidths=[2 * inch] * 3)
    est = estilo('#1e40af')
    est.add('FONTSIZE', (0, 1), (-1, 1), 13)
    est.add('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold')
    est.add('ALIGN', (0, 0), (-1, -1), 'CENTER')
    est.add('TEXTCOLOR', (2, 1), (2, 1), colors.HexColor('#059669' if resumen['balance'] >= 0 else '#dc2626'))
    totales.setStyle(est)
    story += [totales, Spacer(1, 0.25 * inch)]

    png = grafico_png(resumen)
    if png:
        img = Image(png)
        ancho = 6.2 * inch
        img.drawHeight = img.drawHeight * ancho / img.drawWidth
        img.drawWidth = ancho
        story += [img, Spacer(1, 0.15 * inch)]

    if resumen['gastos_por_categoria']:
        story.append(Paragraph("Gastos por categoría", h2))
        data = [['Categoría', '%', 'Monto']]
        for cat, monto in resumen['gastos_por_categoria']:
            pct = 100 * monto / resumen['total_gastos'] if resumen['total_gastos'] else 0
            data.append([cat.replace('_', ' '), f"{pct:.0f}%", dinero(monto)])
        data.append(['TOTAL', '', dinero(resumen['total_gastos'])])
        t = Table(data, colWidths=[3 * inch, 1 * inch, 1.6 * inch])
        e = estilo('#059669')
        e.add('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold')
        t.setStyle(e)
        story += [t, Spacer(1, 0.2 * inch)]

    if resumen['semanas']:
        story.append(Paragraph("Por semana (lunes a domingo)", h2))
        data = [['Semana', 'Ingresos', 'Gastos', 'Balance']]
        for s in resumen['semanas']:
            data.append([f"{s['lunes'].strftime('%d/%m')} – {s['domingo'].strftime('%d/%m/%Y')}",
                         dinero(s['ingresos']), dinero(s['gastos']), dinero(s['balance'])])
        t = Table(data, colWidths=[2.2 * inch, 1.3 * inch, 1.3 * inch, 1.3 * inch])
        t.setStyle(estilo('#1e40af'))
        story += [t, Spacer(1, 0.2 * inch)]

    if resumen['movimientos']:
        story.append(Paragraph(f"Movimientos ({len(resumen['movimientos'])})", h2))
        data = [['Fecha', 'Tipo', 'Categoría', 'Descripción', 'Monto']]
        for m in resumen['movimientos']:
            data.append([m['fecha'].strftime('%d/%m/%Y'), m['tipo'], m['categoria'].replace('_', ' '),
                         str(m['descripcion'])[:32], dinero(m['monto'])])
        t = Table(data, colWidths=[0.9 * inch, 0.7 * inch, 1.1 * inch, 2.5 * inch, 1.0 * inch], repeatRows=1)
        t.setStyle(estilo('#374151'))
        story.append(t)
    else:
        story.append(Paragraph("No hay movimientos en este período.", styles['Normal']))

    SimpleDocTemplate(ruta, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch).build(story)
    return ruta
